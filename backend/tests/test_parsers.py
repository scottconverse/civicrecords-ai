import tempfile
import zipfile
from pathlib import Path
from app.ingestion.parsers import detect_parser, is_image_file
from app.ingestion.parsers.text import TextParser
from app.ingestion.parsers.csv_parser import CsvParser
from app.ingestion.parsers.html import HtmlParser
from app.ingestion.parsers.docx import DocxParser
from app.ingestion.parsers.xlsx import XlsxParser

def test_text_parser():
    with tempfile.NamedTemporaryFile(suffix=".txt", mode="w", delete=False, encoding="utf-8") as f:
        f.write("Hello world\n\nThis is a test document.")
        f.flush()
        parser = TextParser()
        result = parser.parse(Path(f.name))
        assert "Hello world" in result.full_text
        assert result.total_chars > 0
        assert len(result.pages) == 1

def test_csv_parser():
    with tempfile.NamedTemporaryFile(suffix=".csv", mode="w", delete=False, encoding="utf-8") as f:
        f.write("name,age,city\nAlice,30,Denver\nBob,25,Boulder\n")
        f.flush()
        parser = CsvParser()
        result = parser.parse(Path(f.name))
        assert "Alice" in result.full_text
        assert "Denver" in result.full_text
        assert result.metadata["row_count"] == 3

def test_html_parser():
    with tempfile.NamedTemporaryFile(suffix=".html", mode="w", delete=False, encoding="utf-8") as f:
        f.write("<html><head><title>Test</title></head><body><p>Hello HTML</p><script>var x=1;</script></body></html>")
        f.flush()
        parser = HtmlParser()
        result = parser.parse(Path(f.name))
        assert "Hello HTML" in result.full_text
        assert "var x=1" not in result.full_text
        assert result.metadata.get("title") == "Test"

def test_detect_parser_txt():
    parser = detect_parser(Path("test.txt"))
    assert parser is not None
    assert isinstance(parser, TextParser)

def test_detect_parser_pdf():
    assert detect_parser(Path("report.pdf")) is not None

def test_detect_parser_unknown():
    assert detect_parser(Path("file.xyz123")) is None

def test_is_image_file():
    assert is_image_file(Path("scan.jpg")) is True
    assert is_image_file(Path("photo.png")) is True
    assert is_image_file(Path("doc.pdf")) is False
    assert is_image_file(Path("file.txt")) is False


def test_xls_legacy_format_blocklisted():
    """Legacy .xls (BIFF8) is not in XlsxParser.supported_extensions."""
    parser = XlsxParser()
    assert not parser.can_parse(Path("legacy.xls"))
    assert parser.can_parse(Path("modern.xlsx"))
    assert parser.can_parse(Path("macro.xlsm"))


# --- DOCX macro stripping tests ---

def _create_docx(path: Path, text: str = "Test paragraph"):
    """Create a minimal valid .docx file using python-docx."""
    from docx import Document as DocxDocument
    doc = DocxDocument()
    doc.add_paragraph(text)
    doc.save(str(path))


def _inject_vba_entry(docx_path: Path, entry_name: str = "word/vbaProject.bin"):
    """Inject a fake VBA entry into an existing .docx ZIP archive."""
    tmp = docx_path.with_suffix(".tmp")
    with zipfile.ZipFile(docx_path, "r") as zin:
        with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
            for item in zin.namelist():
                zout.writestr(item, zin.read(item))
            zout.writestr(entry_name, b"FAKE_VBA_MACRO_CONTENT")
    tmp.replace(docx_path)


def test_docx_parser_basic():
    """DocxParser extracts text from a clean .docx file."""
    with tempfile.TemporaryDirectory() as td:
        path = Path(td) / "test.docx"
        _create_docx(path, "Hello from DOCX parser test")
        parser = DocxParser()
        result = parser.parse(path)
        assert "Hello from DOCX parser test" in result.full_text
        assert result.metadata["paragraph_count"] >= 1
        assert "macros_stripped" not in result.metadata


def test_docx_macro_stripping():
    """DocxParser strips VBA macros and sets metadata flag."""
    with tempfile.TemporaryDirectory() as td:
        path = Path(td) / "macro.docx"
        _create_docx(path, "Content with macros")
        _inject_vba_entry(path, "word/vbaProject.bin")
        # Confirm the VBA entry is in the file
        with zipfile.ZipFile(path) as z:
            assert "word/vbaProject.bin" in z.namelist()
        parser = DocxParser()
        result = parser.parse(path)
        assert "Content with macros" in result.full_text
        assert result.metadata["macros_stripped"] is True
        assert "word/vbaProject.bin" in result.metadata["stripped_entries"]


def test_docx_clean_file_no_stripping():
    """Clean .docx has no macros_stripped metadata."""
    with tempfile.TemporaryDirectory() as td:
        path = Path(td) / "clean.docx"
        _create_docx(path, "Clean document")
        parser = DocxParser()
        result = parser.parse(path)
        assert result.metadata.get("macros_stripped") is None


# --- XLSX macro stripping tests ---

def _create_xlsx(path: Path, data: list[list[str]] | None = None):
    """Create a minimal valid .xlsx file using openpyxl."""
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    for row in (data or [["Name", "Value"], ["Test", "123"]]):
        ws.append(row)
    wb.save(str(path))


def _inject_xlsx_vba(xlsx_path: Path):
    """Inject a fake VBA entry into an existing .xlsx ZIP archive."""
    tmp = xlsx_path.with_suffix(".tmp")
    with zipfile.ZipFile(xlsx_path, "r") as zin:
        with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
            for item in zin.namelist():
                zout.writestr(item, zin.read(item))
            zout.writestr("xl/vbaProject.bin", b"FAKE_VBA_MACRO_CONTENT")
    tmp.replace(xlsx_path)


def test_xlsx_parser_basic():
    """XlsxParser extracts text from a clean .xlsx file."""
    with tempfile.TemporaryDirectory() as td:
        path = Path(td) / "test.xlsx"
        _create_xlsx(path, [["City", "State"], ["Denver", "CO"]])
        parser = XlsxParser()
        result = parser.parse(path)
        assert "Denver" in result.full_text
        assert result.metadata["sheet_count"] == 1
        assert "macros_stripped" not in result.metadata


def test_xlsx_macro_stripping():
    """XlsxParser strips VBA macros and sets metadata flag."""
    with tempfile.TemporaryDirectory() as td:
        path = Path(td) / "macro.xlsx"
        _create_xlsx(path, [["Data", "Here"]])
        _inject_xlsx_vba(path)
        with zipfile.ZipFile(path) as z:
            assert "xl/vbaProject.bin" in z.namelist()
        parser = XlsxParser()
        result = parser.parse(path)
        assert "Data" in result.full_text
        assert result.metadata["macros_stripped"] is True
        assert "xl/vbaProject.bin" in result.metadata["stripped_entries"]
